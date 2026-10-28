"""FlexiCred — painel de prospecção. Inicie com: streamlit run 2_app.py"""
import io
import re
import json
import base64
from pathlib import Path

import duckdb
import pandas as pd
import streamlit as st

RAIZ = Path(__file__).resolve().parent
st.set_page_config(page_title="FlexiCred | Prospecção", page_icon="🔎", layout="wide")
st.markdown((RAIZ / "assets" / "estilo.css").read_text(encoding="utf-8"), unsafe_allow_html=True)
logo = base64.b64encode((RAIZ / "assets" / "logo_flexi.png").read_bytes()).decode()
st.markdown(f"""<div class="brandbar">
<div class="brand-logo" style="background-image:url(data:image/png;base64,{logo})" role="img" aria-label="Grupo Flexi"></div>
<div class="brand-divider"></div><span>FLEXICRED <b>/</b> PROSPECÇÃO EMPRESARIAL</span>
<span class="brand-tag">INTELIGÊNCIA COMERCIAL</span></div>""", unsafe_allow_html=True)
st.title("Prospecção de empresas")
st.caption("Encontre empresas por região, atividade econômica e perfil de negócio.")

import hashlib
import logging
import os
import shutil
import tempfile
import threading
import time
import urllib.request
import zipfile
from pathlib import Path

URL_BASE = "https://github.com/guimaraesdherick-create/prospeccao-empresas/releases/download/dados-2026-10/prospeccao.zip"
SHA256_ZIP = "31ddea72b88d08af8586bde4c523e489375b02a2b668c4320c6ac5ce1410009f"
CONFIG_DB = {"memory_limit": "384MB", "threads": "1"}


@st.cache_resource
def trava_preparacao():
    return threading.Lock()


@st.cache_resource(show_spinner=False)
def preparar_base(url, checksum):
    # Preserva o uso local da base existente.
    local = Path(__file__).resolve().parent / "prospeccao.duckdb"
    if local.is_file():
        return str(local)

    pasta = Path(tempfile.gettempdir()) / "prospeccao" / checksum
    pasta.mkdir(parents=True, exist_ok=True)
    destino = pasta / "prospeccao.duckdb"

    with trava_preparacao():
        if destino.is_file():
            return str(destino)

        # ZIP + banco extraído coexistem durante a preparação.
        if shutil.disk_usage(pasta).free < 6 * 1024**3:
            raise RuntimeError("São necessários pelo menos 6 GiB livres para preparar a base.")

        with tempfile.TemporaryDirectory(dir=pasta) as temporario:
            zip_path = Path(temporario) / "base.zip"
            banco_path = Path(temporario) / "base.duckdb"
            status = st.empty()
            try:
                for tentativa in range(3):
                    try:
                        digest = hashlib.sha256()
                        requisicao = urllib.request.Request(
                            url, headers={"User-Agent": "prospeccao-streamlit"}
                        )
                        with urllib.request.urlopen(requisicao, timeout=120) as resposta:
                            total = int(resposta.headers.get("Content-Length") or 0)
                            baixados = 0
                            with zip_path.open("wb") as arquivo:
                                while True:
                                    bloco = resposta.read(4 * 1024 * 1024)
                                    if not bloco:
                                        break
                                    arquivo.write(bloco)
                                    digest.update(bloco)
                                    baixados += len(bloco)
                                    status.info(
                                        f"Baixando base: {baixados / 1024**2:.0f} MB"
                                        + (f" de {total / 1024**2:.0f} MB" if total else "")
                                    )
                        if digest.hexdigest() != checksum:
                            raise ValueError("O SHA-256 do download não confere.")
                        break
                    except Exception:
                        if tentativa == 2:
                            raise
                        status.info("Download interrompido. Tentando novamente...")
                        time.sleep(2)

                status.info("Download verificado. Extraindo a base...")
                with zipfile.ZipFile(zip_path) as pacote:
                    candidatos = [
                        item for item in pacote.infolist()
                        if not item.is_dir()
                        and item.filename.replace("\\", "/").split("/")[-1]
                        == "prospeccao.duckdb"
                    ]
                    if len(candidatos) != 1:
                        raise ValueError(
                            "O ZIP deve conter exatamente um arquivo prospeccao.duckdb."
                        )
                    # Copia somente o banco, sem extrair caminhos fornecidos pelo ZIP.
                    with pacote.open(candidatos[0]) as origem, banco_path.open("wb") as saida:
                        shutil.copyfileobj(origem, saida, length=4 * 1024 * 1024)

                with duckdb.connect(
                    str(banco_path), read_only=True, config=CONFIG_DB
                ) as teste:
                    teste.execute("SELECT * FROM base LIMIT 0")
                os.replace(banco_path, destino)
            finally:
                status.empty()

        return str(destino)


try:
    with st.spinner("Preparando a base. Na primeira execução, aguarde alguns minutos..."):
        DB = preparar_base(URL_BASE, SHA256_ZIP)
except Exception as erro:
    logging.exception("Falha na preparação da base")
    st.error(f"Não foi possível preparar a base: {erro}")
    st.stop()


@st.cache_data
def catalogo_cnaes():
    return json.loads((RAIZ / "assets" / "cnaes.json").read_text(encoding="utf-8"))


def conectar():
    con = duckdb.connect(DB, read_only=True, config=CONFIG_DB)
    try:
        catalogo = pd.DataFrame(list(catalogo_cnaes().items()), columns=["codigo", "descricao"])
        con.register("catalogo_cnae", catalogo)
        con.execute("""CREATE TEMP VIEW base_consulta AS
            SELECT b.* REPLACE (
                coalesce(c.descricao,
                    CASE WHEN lower(trim(coalesce(b.cnae_descricao, ''))) NOT IN ('', 'none', 'null', 'nan')
                         THEN b.cnae_descricao END,
                    'Descrição não disponível no catálogo') AS cnae_descricao
            ) FROM main.base b
            LEFT JOIN catalogo_cnae c ON trim(b.cnae_principal) = c.codigo
        """)
        return con
    except Exception:
        con.close()
        raise




@st.cache_data
def listar_ufs():
    with conectar() as con:
        return [r[0] for r in con.execute(
            "SELECT DISTINCT uf FROM base_consulta WHERE uf IS NOT NULL ORDER BY uf").fetchall()]


@st.cache_data
def listar_municipios(ufs):
    if not ufs:
        return []
    with conectar() as con:
        return [r[0] for r in con.execute(
            "SELECT DISTINCT municipio FROM base_consulta WHERE uf IN (SELECT unnest(?::VARCHAR[])) "
            "AND municipio IS NOT NULL ORDER BY municipio", [list(ufs)]).fetchall()]


@st.cache_data
def listar_cnaes():
    with conectar() as con:
        rows = con.execute(
            "SELECT cnae_principal, any_value(cnae_descricao), count(*) "
            "FROM base_consulta WHERE cnae_principal IS NOT NULL "
            "GROUP BY 1 ORDER BY 1").fetchall()
    return rows


def fmt_cnae(c):
    return f"{c[:4]}-{c[4]}/{c[5:]}" if c and len(c) == 7 else c


def so_digitos(txt):
    return re.sub(r"\D", "", txt)


# ---------------------------------------------------------------- filtros
try:
    ufs_disp = listar_ufs()
except Exception as e:  # base não existe
    logging.exception("Falha na consulta da base")
    st.error(f"Não consegui consultar a base: {e}")
    st.stop()

cnaes = listar_cnaes()
opcoes_cnae = {f"{fmt_cnae(str(c))} · {d} ({n:,})": str(c) for c, d, n in cnaes}

with st.sidebar:
    st.markdown('<div class="side-kicker">SEGMENTAÇÃO</div>', unsafe_allow_html=True)
    st.header("Filtros de busca")
    st.caption("Escolha a região e refine seu público.")
    st.subheader("Localização")
    ufs = st.multiselect("Estados", ufs_disp, default=["SC"] if "SC" in ufs_disp else [], placeholder="Selecione os estados")
    municipios = st.multiselect("Municípios", listar_municipios(tuple(ufs)), placeholder="Todos dos estados selecionados")
    buscar = st.button("Buscar empresas", type="primary", width="stretch")

    with st.expander("Atividade econômica", expanded=True):
        escolhidos = st.multiselect("CNAE principal", list(opcoes_cnae), placeholder="Busque por código ou atividade")
        colados = st.text_area("Códigos CNAE", placeholder="4520-0/01, 9602501", height=80, help="Separe os códigos por vírgula, ponto e vírgula ou linha.")
        incluir_sec = st.checkbox("Considerar CNAEs secundários")
    with st.expander("Perfil e endereço"):
        bairro_txt = st.text_input("Bairro contém")
        portes = st.multiselect("Porte", ["Microempresa", "Empresa de pequeno porte", "Demais", "Não informado"], placeholder="Todos os portes")
        idade_min = st.slider("Tempo mínimo de atividade (anos)", 0, 30, 0)
        so_matriz = st.checkbox("Somente matriz", value=True)
    with st.expander("Dados de contato"):
        exige_email = st.checkbox("Com e-mail")
        exige_tel = st.checkbox("Com telefone")
        ocultar_contab = st.checkbox("Ocultar e-mails de contabilidade", value=True, help="Filtro por termos no endereço de e-mail; não confirma a titularidade do contato.")
    st.caption("Receita Federal · Dados Abertos CNPJ")

# ---------------------------------------------------------------- consulta
codigos = {opcoes_cnae[k] for k in escolhidos}
invalidos = []
for tok in re.split(r"[,\n;]+", colados):
    if not tok.strip():
        continue
    d = so_digitos(tok)
    if len(d) == 7:
        codigos.add(d)
    else:
        invalidos.append(tok.strip())
if invalidos:
    st.warning("Revise os códigos CNAE: cada código deve ter 7 dígitos. Valores inválidos: " + ", ".join(invalidos))
    st.stop()

where, params = [], []

if ufs:
    where.append("uf IN (SELECT unnest(?::VARCHAR[]))")
    params.append(list(ufs))
if municipios:
    where.append("municipio IN (SELECT unnest(?::VARCHAR[]))")
    params.append(list(municipios))
if bairro_txt.strip():
    where.append("bairro ILIKE ?")
    params.append(f"%{bairro_txt.strip()}%")

if codigos:
    lista = sorted(codigos)
    cond = "cnae_principal IN (SELECT unnest(?::VARCHAR[]))"
    params.append(lista)
    if incluir_sec:
        cond += (" OR list_has_any(string_split(coalesce(cnae_secundaria,''), ','), ?::VARCHAR[])")
        params.append(lista)
    where.append(f"({cond})")

if portes:
    where.append("porte IN (SELECT unnest(?::VARCHAR[]))")
    params.append(portes)
if idade_min:
    where.append("data_inicio <= current_date - to_years(?::INTEGER)")
    params.append(idade_min)
if so_matriz:
    where.append("eh_matriz")
if exige_email:
    where.append("email IS NOT NULL")
if exige_tel:
    where.append("telefone IS NOT NULL")
if ocultar_contab:
    where.append("(email IS NULL OR email NOT ILIKE '%contab%' "
                 "AND email NOT ILIKE '%contador%' AND email NOT ILIKE '%assessoria%')")

sql_where = ("WHERE " + " AND ".join(where)) if where else ""

if not codigos and not ufs:
    st.info("Escolha ao menos um estado ou um CNAE na barra lateral para começar.")
    st.stop()

assinatura = repr((sql_where, params))
if buscar:
    st.session_state["filtros_aplicados"] = assinatura
if st.session_state.get("filtros_aplicados") != assinatura:
    st.markdown("""<div class="empty-panel"><span class="empty-label">SUA PRÓXIMA OPORTUNIDADE COMEÇA AQUI</span>
    <h3>Defina o perfil das empresas que deseja encontrar</h3>
    <p>Selecione os estados, procure uma atividade econômica e clique em <strong>Buscar empresas</strong>.</p>
    <div class="steps"><span>01 · Selecione a região</span><span>02 · Refine a atividade</span><span>03 · Exporte sua lista</span></div></div>""", unsafe_allow_html=True)
    st.stop()

with conectar() as con:
    total = con.execute(f"SELECT count(*) FROM base_consulta {sql_where}", params).fetchone()[0]
c1, c2, c3 = st.columns(3)
c1.metric("Empresas encontradas", f"{total:,}")
c2.metric("CNAEs selecionados", len(codigos) or "todos")
c3.metric("Municípios", len(municipios) or ("todos" if ufs else "—"))

if total == 0:
    st.warning("Nenhuma empresa com esses filtros.")
    st.stop()

LIMITE_EXPORT = 20_000
with conectar() as con:
    df = con.execute(
        f"SELECT * FROM base_consulta {sql_where} ORDER BY municipio, razao_social LIMIT {LIMITE_EXPORT}",
        params).fetchdf()

df["cnae_principal"] = df["cnae_principal"].map(fmt_cnae)


def link_whatsapp(tel):
    if not isinstance(tel, str) or len(tel) < 10:
        return None
    numero = tel[2:]
    if len(numero) == 9 and numero.startswith("9"):  # celular
        return f"https://wa.me/55{tel}"
    return None


df["whatsapp"] = df["telefone"].map(link_whatsapp)
df = df.drop(columns=["eh_matriz"])

st.subheader("Empresas selecionadas")
st.caption("Confira os contatos e exporte sua lista para organizar o atendimento.")

visiveis = [c for c in ["razao_social", "nome_fantasia", "cnpj", "municipio", "uf", "cnae_principal", "cnae_descricao", "telefone", "email", "whatsapp", "porte", "data_inicio", "capital_social"] if c in df]
st.dataframe(df[visiveis].head(1000), width="stretch", height=460,
             column_config={
                 "cnpj": st.column_config.TextColumn("CNPJ"),
                 "razao_social": st.column_config.TextColumn("Razão social"),
                 "nome_fantasia": st.column_config.TextColumn("Nome fantasia"),
                 "capital_social": st.column_config.NumberColumn("Capital social (R$)", format="%.2f"),
                 "cnae_principal": st.column_config.TextColumn("CNAE principal"),
                 "cnae_descricao": st.column_config.TextColumn("Atividade"),
                 "municipio": st.column_config.TextColumn("Município"),
                 "telefone": st.column_config.TextColumn("Telefone"),
                 "email": st.column_config.TextColumn("E-mail"),
                 "data_inicio": st.column_config.DateColumn("Abertura", format="DD/MM/YYYY"),
                 "whatsapp": st.column_config.LinkColumn("WhatsApp", display_text="Abrir conversa"),
             }, hide_index=True)
if total > 1000:
    st.caption(f"Mostrando as 1.000 primeiras linhas; a planilha exportada traz "
               f"até {LIMITE_EXPORT:,}.")

st.divider()
st.subheader("Exporte sua lista")
st.caption("Até 20.000 empresas por exportação, respeitando os filtros selecionados.")

# Os arquivos são gerados somente quando solicitados.
if st.session_state.get("export_assinatura") != assinatura:
    st.session_state.pop("export_arquivos", None)
    st.session_state.pop("export_assinatura", None)

if st.button("Preparar Excel e CSV", width="content"):
    with st.spinner("Preparando arquivos de exportação..."):
        buf = io.BytesIO()
        # Grava strings como texto no Excel, inclusive valores que começam com '='.
        with pd.ExcelWriter(buf, engine="openpyxl") as w:
            df.to_excel(w, index=False, sheet_name="Prospects")
            ws = w.sheets["Prospects"]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            from openpyxl.styles import Font, PatternFill
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="101A30")
                cell.font = Font(color="FFFFFF", bold=True)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str) and cell.data_type == "f":
                        cell.data_type = "s"
            for col in ws.columns:
                tamanho = max(len(str(c.value)) if c.value is not None else 0 for c in col[:200])
                ws.column_dimensions[col[0].column_letter].width = min(max(tamanho + 2, 12), 48)
        csv_df = df.copy()
        for col in csv_df.select_dtypes(include=["object", "string"]).columns:
            csv_df[col] = csv_df[col].map(lambda x: "'" + x if isinstance(x, str) and x.lstrip().startswith(("=", "+", "-", "@")) else x)
        st.session_state["export_arquivos"] = (buf.getvalue(), csv_df.to_csv(index=False, sep=";").encode("utf-8-sig"))
        st.session_state["export_assinatura"] = assinatura

if "export_arquivos" in st.session_state:
    excel, csv_bytes = st.session_state["export_arquivos"]
    b1, b2, _ = st.columns([1, 1, 3])
    b1.download_button("Baixar Excel", excel, "flexicred_prospects.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", on_click="ignore", width="stretch")
    b2.download_button("Baixar CSV", csv_bytes, "flexicred_prospects.csv", "text/csv", on_click="ignore", width="stretch")
st.caption("FlexiCred · Informações cadastrais para apoio à prospecção. A base disponível representa os arquivos importados.")
